import io
import os
from pathlib import Path
from typing import Any
from fastapi import Body, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from .analyzer import analyze, train_model

app = FastAPI(title='CourtInsight API', version='0.1.0')
default_origins = [
    'http://localhost:5173',
    'http://localhost:5174',
    'http://localhost:5175',
    'http://localhost:5176',
    'http://localhost:5177',
    'http://127.0.0.1:5173',
    'http://127.0.0.1:5174',
    'http://127.0.0.1:5175',
    'http://127.0.0.1:5176',
    'http://127.0.0.1:5177',
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=[x.strip() for x in os.getenv('CORS_ORIGINS', ','.join(default_origins)).split(',') if x.strip()],
    allow_origin_regex=r'https?://(localhost|127\.0\.0\.1):\d+',
    allow_credentials=False,
    allow_methods=['GET', 'POST', 'OPTIONS'],
    allow_headers=['*'],
)
MAX_BYTES = 10 * 1024 * 1024

class TextInput(BaseModel):
    text: str = Field(min_length=40, max_length=300_000)
    use_embeddings: bool = False

class TrainingItem(BaseModel):
    text: str = Field(min_length=1)
    label: str

class TrainingRequest(BaseModel):
    examples: list[TrainingItem]


def normalize_training_examples(payload: Any):
    if isinstance(payload, list):
        examples = payload
    elif isinstance(payload, dict):
        examples = payload.get('examples') or payload.get('data') or payload.get('items') or []
    else:
        raise HTTPException(status_code=422, detail='Training payload must be a list of examples or an object with an examples list.')

    if not isinstance(examples, list):
        raise HTTPException(status_code=422, detail='Training examples must be provided as a list.')

    rows = []
    for index, item in enumerate(examples, start=1):
        if not isinstance(item, dict):
            raise HTTPException(status_code=422, detail=f'Example {index} must be an object with text and label.')

        raw_text = str(item.get('text', '')).strip()
        raw_label = str(item.get('label', '')).strip()
        if not raw_text:
            raise HTTPException(status_code=422, detail=f'Example {index} is missing text.')

        normalized_label = {'good': 'Good', 'better': 'Better', 'worst': 'Worst'}.get(raw_label.lower(), raw_label)
        if normalized_label not in {'Good', 'Better', 'Worst'}:
            raise HTTPException(status_code=422, detail='Labels must be Good, Better, or Worst.')

        rows.append({'text': raw_text, 'label': normalized_label})

    return rows

@app.get('/api/health')
def health():
    return {'status': 'ok', 'service': 'CourtInsight'}

@app.post('/api/analyze/text')
def analyze_text(body: TextInput):
    try:
        return analyze(body.text, body.use_embeddings)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

@app.post('/api/analyze/file')
async def analyze_file(file: UploadFile = File(...), use_embeddings: bool = Form(False)):
    suffix = Path(file.filename or '').suffix.lower()
    if suffix not in {'.pdf', '.docx', '.txt'}:
        raise HTTPException(status_code=415, detail='Only PDF, DOCX and TXT files are supported.')
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail='File exceeds 10 MB.')
    try:
        if suffix == '.pdf':
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError('Password-protected PDFs are not supported.')
            if len(reader.pages) > 100:
                raise ValueError('PDF limit is 100 pages.')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        elif suffix == '.docx':
            doc = Document(io.BytesIO(data))
            text = '\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text += '\n' + '\n'.join(' '.join(cell.text for cell in row.cells) for row in table.rows)
        else:
            text = data.decode('utf-8-sig')
        if not text.strip():
            raise ValueError('No text found. Scanned PDFs require OCR, which is not included.')
        if len(text) > 300_000:
            raise ValueError('Extracted text exceeds the 300,000-character limit.')
        return {'filename': file.filename, **analyze(text, use_embeddings)}
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f'Unable to process document ({type(exc).__name__}).') from exc

@app.post('/api/train')
def train(payload: Any = Body(...)):
    try:
        rows = normalize_training_examples(payload)
    except HTTPException:
        raise

    try:
        count = train_model(rows)
        return {'trained': True, 'examples': count, 'note': 'Model is in-memory and resets on server restart. Evaluate with held-out labeled legal texts before practical use.'}
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
